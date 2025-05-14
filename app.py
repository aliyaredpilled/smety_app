#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Flask‑приложение: прогресс‑бар 0‑100 %, декод русских имён в ZIP (включая пути),
улучшенный отчет (обработано / пустые / ошибки + списки), удаление корневой папки из пути.
"""
import os
import re
import uuid
import zipfile
import shutil
import traceback
import logging
import openpyxl

import dispatcher
from flask import (
    Flask, request, render_template, jsonify,
    send_from_directory, url_for
)
from werkzeug.utils import secure_filename
from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter
from formatting import apply_reference_widths, auto_adjust_column_width, apply_formatting

# ────────── базовая конфигурация ──────────
logging.basicConfig(level=logging.INFO, format="[%(asctime)s] %(levelname)s: %(message)s")
app = Flask(__name__)
app.config.update(SEND_FILE_MAX_AGE_DEFAULT=0, TEMPLATES_AUTO_RELOAD=True, DEBUG=True)

UPLOAD_FOLDER, RESULTS_FOLDER, REF_FOLDER = "uploads", "results", "reference_files"
ALLOWED_EXT = {"xlsx", "xlsm", "zip"}

for p in (UPLOAD_FOLDER, RESULTS_FOLDER, REF_FOLDER):
    os.makedirs(p, exist_ok=True)

REF_SMETA_RU   = os.path.join(REF_FOLDER, "Смета ру.xlsm")
REF_TURBO      = os.path.join(REF_FOLDER, "Турбосметчик1,2,3.xlsm")
REF_GRAND      = os.path.join(REF_FOLDER, "Пример1 2.xlsx")

processing_status: dict[str, dict] = {}

# ────────── утилиты ──────────
def allowed_file(name: str) -> bool:
    """Проверяет, имеет ли файл разрешенное расширение."""
    return "." in name and name.rsplit(".", 1)[1].lower() in ALLOWED_EXT

def calc_percent(stage: str, done: int, total: int) -> int:
    """Рассчитывает процент выполнения для прогресс-бара."""
    if stage in ("upload", "extract"):
        return 0 if stage == "upload" else 5
    if stage == "processing":
        return 5 + round(90 * done / total) if total else 5
    if stage == "finalize":
        return 95
    if stage == "done":
        return 100
    return 0

# ────────── routes ──────────
@app.route("/")
def index():
    """Отображает главную страницу."""
    try:
        types = dispatcher.get_available_processor_types()
        return render_template("index.html", smeta_types=types)
    except Exception:
        logging.exception("Не удалось загрузить типы смет")
        return render_template("index.html", smeta_types=[], error="Ошибка загрузки типов смет")

@app.route("/upload", methods=["POST"])
def upload_file():
    """Обрабатывает загрузку файла, извлечение (если ZIP) и запуск обработки."""
    sess = request.form.get("client_session_id") or str(uuid.uuid4())
    f    = request.files.get("file")
    smeta_type = request.form.get("smeta_type")

    if not f or f.filename == "": return jsonify(success=False, error="Файл не выбран"), 400
    if not smeta_type: return jsonify(success=False, error="Тип сметы не выбран"), 400
    if not allowed_file(f.filename): return jsonify(success=False, error="Неподдерживаемый тип файла (разрешены: xlsx, xlsm, zip)"), 400

    # Инициализируем статус сессии как можно раньше
    processing_status[sess] = { 
        "processed": 0, 
        "total": 0, # Общее количество файлов будет обновлено позже, если это ZIP
        "status": "Загрузка файла...", # Начальный статус
        "error": None, 
        "percent": calc_percent("upload", 0, 0) 
    }

    work_dir = os.path.join(UPLOAD_FOLDER, sess)
    try:
        logging.info(f"({sess}) Проверка существования рабочей папки (перед созданием новой): {work_dir}")
        path_exists = os.path.exists(work_dir)
        logging.info(f"({sess}) Результат проверки os.path.exists({work_dir}): {path_exists}")
        if path_exists:
            logging.info(f"({sess}) Попытка удаления рабочей папки (перед созданием новой): {work_dir}")
            try:
                shutil.rmtree(work_dir)
                logging.info(f"({sess}) Рабочая папка {work_dir} успешно удалена (перед созданием новой).")
            except Exception as e:
                logging.error(f"({sess}) Ошибка при удалении рабочей папки {work_dir} (перед созданием новой): {e}")
                # Решаем, нужно ли прерывать операцию, если папка не удалилась.
                # В данном случае, создание новой папки может перезаписать старую,
                # но если там остались файлы, которые не перезапишутся, могут быть проблемы.
                # Пока оставим как есть, но это место для возможного улучшения.
        os.makedirs(work_dir, exist_ok=True)
    except OSError as e:
        logging.error(f"({sess}) Не удалось создать рабочую папку {work_dir}: {e}")
        # Обновляем статус на ошибку, если не удалось создать папку
        if sess in processing_status: # Сессия должна была быть инициализирована выше
            processing_status[sess].update(status="Ошибка сервера", error=f"Ошибка файловой системы: {e}", percent=100)
        return jsonify(success=False, error="Ошибка файловой системы на сервере."), 500

    orig_filename = f.filename # <<< Сохраняем исходное имя файла
    logging.info(f"({sess}) Получено имя файла от клиента: '{orig_filename}'")
    is_zip = orig_filename.lower().endswith(".zip")

    _secure_attempt = secure_filename(orig_filename)

    _orig_basename, _orig_extension = os.path.splitext(orig_filename)
    _orig_extension_clean = _orig_extension.lower().lstrip('.')

    # Проверяем, если secure_filename вернул пустоту или только расширение (и было имя файла)
    if not _secure_attempt or (_secure_attempt.lower() == _orig_extension_clean and _orig_basename):
        if _orig_extension_clean in ALLOWED_EXT:
            safe_filename_base = f"uploaded.{_orig_extension_clean}"
        else: # Запасной вариант, если оригинальное расширение не разрешено (хотя это должно быть поймано раньше)
            safe_filename_base = "uploaded.zip" if is_zip else "uploaded.xlsx"
    else:
        # secure_filename вернул что-то, что похоже на полноценное имя
        safe_filename_base = _secure_attempt

    # Дополнительный запасной вариант, если safe_filename_base все еще пуст
    if not safe_filename_base:
        safe_filename_base = "uploaded.zip" if is_zip else "uploaded.xlsx"

    logging.info(f"({sess}) Имя файла для сохранения на диске (после обработки): '{safe_filename_base}'")
    saved_filepath = os.path.join(work_dir, safe_filename_base)

    try:
        f.save(saved_filepath)
        logging.info(f"({sess}) Файл '{orig_filename}' сохранен как '{saved_filepath}'")
    except Exception as e:
        logging.error(f"({sess}) Не удалось сохранить файл {orig_filename}: {e}")
        processing_status[sess].update(status="Ошибка сохранения", error=str(e))
        return jsonify(success=False, error=f"Ошибка сохранения файла: {e}"), 500

    files_to_proc = []
    try:
        if is_zip:
            # --- Обработка ZIP архива ---
            processing_status[sess].update(status="Распаковка архива...", percent=calc_percent("extract", 0, 0))
            extract_path = os.path.join(work_dir, "extracted")
            os.makedirs(extract_path, exist_ok=True)
            processed_zip_paths = set()

            # --- ИЗМЕНЕНИЕ: Получаем базовое имя ZIP файла ---
            zip_basename = os.path.splitext(orig_filename)[0]
            logging.debug(f"({sess}) Базовое имя ZIP файла: '{zip_basename}'")
            # --- КОНЕЦ ИЗМЕНЕНИЯ ---

            try:
                with zipfile.ZipFile(saved_filepath, 'r') as zip_ref:
                    members_to_extract = [
                        m for m in zip_ref.infolist()
                        if not m.is_dir()
                        and not m.filename.startswith(('__MACOSX/', '.'))
                        and not os.path.basename(m.filename).startswith(('._', '~$', '$'))
                        and m.filename.lower().endswith(('.xlsx', '.xlsm'))
                    ]
                    processing_status[sess]["total"] = len(members_to_extract)
                    logging.info(f"({sess}) Найдено файлов в ZIP для обработки: {len(members_to_extract)}")
                    if not members_to_extract: raise ValueError("В архиве не найдено поддерживаемых Excel файлов (.xlsx, .xlsm).")

                    for member in members_to_extract:
                        original_zip_path = member.filename
                        # --- Отладка: Показываем исходное имя файла --- 
                        logging.debug(f"  ({sess}) ZIP Member Original Path: '{original_zip_path}'")
                        # Попытка показать байты (может помочь увидеть "странные" символы)
                        try:
                             logging.debug(f"  ({sess}) ZIP Member Original Path (bytes surrogateescaped): {original_zip_path.encode('utf-8', 'surrogateescape')}")
                        except Exception: pass # Игнорируем ошибки, если имя не кодируется в UTF-8
                        # --- Конец отладки --- 

                        decoded_full_path = None
                        try:
                            decoded_full_path = original_zip_path.encode('cp437').decode('utf-8', 'ignore')
                            # --- Отладка: Результат декодирования UTF-8 ---
                            logging.debug(f"  ({sess}) Декодировано (cp437->utf-8) '{original_zip_path}' -> '{decoded_full_path}'")
                            # --- Конец отладки --- 
                        except UnicodeDecodeError:
                            try:
                                decoded_full_path = original_zip_path.encode('cp437').decode('cp866', 'ignore')
                                # --- Отладка: Результат декодирования CP866 ---
                                logging.warning(f"  ({sess}) Fallback (cp437->cp866): '{original_zip_path}' -> '{decoded_full_path}'")
                                # --- Конец отладки --- 
                            except Exception:
                                # --- Отладка: Ошибка декодирования ---
                                decoded_full_path = original_zip_path
                                logging.error(f"  ({sess}) НЕ УДАЛОСЬ ДЕКОДИРОВАТЬ путь '{original_zip_path}'. Используется как есть.")
                                # --- Конец отладки --- 
                        except Exception as decode_err:
                            # --- Отладка: Ошибка декодирования ---
                            decoded_full_path = original_zip_path
                            logging.error(f"  ({sess}) Ошибка декодирования '{original_zip_path}': {decode_err}. Используется как есть.")
                            # --- Конец отладки --- 

                        normalized_path = decoded_full_path.replace('\\', '/')
                        # --- Отладка: Нормализованный путь ---
                        logging.debug(f"  ({sess}) Normalized Path: '{normalized_path}'")
                        # --- Конец отладки --- 

                        if normalized_path.startswith('/') or '..' in normalized_path.split('/') or re.match(r"^[a-zA-Z]:/", normalized_path):
                            logging.warning(f"({sess}) Пропуск потенциально небезопасного пути: '{decoded_full_path}'")
                            continue

                        sanitized_relative_path = re.sub(r'[<>:"/\\|?*\x00-\x1f]', '_', normalized_path)
                        if sanitized_relative_path in processed_zip_paths:
                            logging.warning(f"({sess}) Пропуск дублирующегося пути после очистки: '{sanitized_relative_path}' (из '{original_zip_path}')")
                            continue
                        processed_zip_paths.add(sanitized_relative_path)
                        if not sanitized_relative_path:
                             logging.warning(f"({sess}) Путь стал пустым после очистки для '{original_zip_path}'. Пропускаем.")
                             continue

                        target_disk_path = os.path.join(extract_path, sanitized_relative_path)

                        # --- ИЗМЕНЕНИЕ: Определяем имя для отображения с удалением корневой папки ---
                        display_name = normalized_path # По умолчанию - нормализованный путь
                        path_parts = normalized_path.split('/')
                        # Проверяем, есть ли папки И совпадает ли первая папка с базовым именем ZIP
                        if len(path_parts) > 1 and path_parts[0] == zip_basename:
                             # Собираем путь заново, начиная со второго элемента
                             display_name = '/'.join(path_parts[1:])
                             logging.debug(f"  ({sess}) Удалена корневая папка '{zip_basename}' для display_name: '{display_name}'")
                        elif len(path_parts) == 1:
                             # Если папок нет, display_name уже равен имени файла
                             pass
                        # Если первая папка не совпадает, оставляем display_name как normalized_path

                        # Если после удаления display_name остался пустым (очень редкий случай), используем basename
                        if not display_name:
                            display_name = os.path.basename(normalized_path)
                        # --- Отладка: Финальный display_name ---
                        logging.debug(f"  ({sess}) Final Display Name: '{display_name}'")
                        # --- Конец отладки --- 
                        # --- КОНЕЦ ИЗМЕНЕНИЯ ---

                        try:
                            os.makedirs(os.path.dirname(target_disk_path), exist_ok=True)
                            with zip_ref.open(member) as source, open(target_disk_path, "wb") as target:
                                shutil.copyfileobj(source, target)

                            # Сохраняем display_name как идентификатор
                            files_to_proc.append({
                                "path": target_disk_path,
                                "original": display_name # Используем скорректированный display_name
                            })
                            logging.info(f"  ({sess}) Успешно извлечено: '{display_name}' -> '{target_disk_path}'")
                        except OSError as ose:
                            logging.error(f"  [ОШИБКА ОС] ({sess}) при извлечении '{display_name}' в '{target_disk_path}': {ose}")
                        except Exception as extract_err:
                            logging.error(f"  [ОШИБКА ИЗВЛЕЧЕНИЯ] ({sess}) для '{display_name}': {extract_err}")

                if not files_to_proc and processing_status[sess]["total"] > 0 :
                     raise ValueError("Не удалось извлечь файлы из архива из-за ошибок.")

                # --- СОРТИРОВКА ФАЙЛОВ ИЗ ZIP ПО ИМЕНИ/ПУТИ --- 
                if is_zip:
                    try:
                        # Функция для естественной сортировки
                        def natural_sort_key(item):
                            filename = item['original']
                            match = re.match(r'^(\d+)\.?.*', filename)
                            if match:
                                # Если нашли число в начале, возвращаем его (int) и полное имя
                                return (int(match.group(1)), filename)
                            else:
                                # Если числа нет, возвращаем inf и полное имя
                                return (float('inf'), filename)

                        files_to_proc.sort(key=natural_sort_key)
                        logging.info(f"({sess}) Файлы из ZIP отсортированы естественным порядком ({len(files_to_proc)} шт.)")
                    except Exception as sort_err:
                         logging.warning(f"({sess}) Не удалось отсортировать файлы из ZIP: {sort_err}")
                # --- КОНЕЦ СОРТИРОВКИ ---

            except zipfile.BadZipFile: raise ValueError("Некорректный или поврежденный ZIP архив.")
            except ValueError as ve: raise ve
            except Exception as e:
                logging.exception(f"({sess}) Непредвиденная ошибка при работе с ZIP")
                raise ValueError(f"Ошибка при распаковке ZIP: {e}")
        else:
            # --- Обработка одиночного Excel файла ---
            processing_status[sess]["total"] = 1
            # Сохраняем только имя файла как идентификатор
            files_to_proc.append({
                "path": saved_filepath,
                "original": os.path.basename(orig_filename)
                })

    except ValueError as e:
        processing_status[sess].update(status="Ошибка подготовки", error=str(e), percent=100)
        logging.error(f"({sess}) Ошибка подготовки файлов: {e}")
        logging.info(f"({sess}) Проверка существования рабочей папки (после ValueError): {work_dir}")
        path_exists = os.path.exists(work_dir)
        logging.info(f"({sess}) Результат проверки os.path.exists({work_dir}) (после ValueError): {path_exists}")
        if path_exists:
            logging.info(f"({sess}) Попытка удаления рабочей папки (после ValueError): {work_dir}")
            try:
                shutil.rmtree(work_dir)
                logging.info(f"({sess}) Рабочая папка {work_dir} успешно удалена (после ValueError).")
            except Exception as e_rm:
                logging.error(f"({sess}) Ошибка при удалении рабочей папки {work_dir} (после ValueError): {e_rm}")
        else:
            logging.info(f"({sess}) Рабочая папка {work_dir} не найдена для удаления (после ValueError).")
        return jsonify(success=False, error=str(e)), 400
    except Exception as e:
        processing_status[sess].update(status="Критическая ошибка", error="Внутренняя ошибка сервера", percent=100)
        logging.exception(f"({sess}) Критическая ошибка на этапе подготовки")
        logging.info(f"({sess}) Проверка существования рабочей папки (после критической ошибки Exception): {work_dir}")
        path_exists = os.path.exists(work_dir)
        logging.info(f"({sess}) Результат проверки os.path.exists({work_dir}) (после крит. ошибки Exception): {path_exists}")
        if path_exists:
            logging.info(f"({sess}) Попытка удаления рабочей папки (после критической ошибки Exception): {work_dir}")
            try:
                shutil.rmtree(work_dir)
                logging.info(f"({sess}) Рабочая папка {work_dir} успешно удалена (после критической ошибки Exception).")
            except Exception as e_rm:
                logging.error(f"({sess}) Ошибка при удалении рабочей папки {work_dir} (после критической ошибки Exception): {e_rm}")
        else:
            logging.info(f"({sess}) Рабочая папка {work_dir} не найдена для удаления (после крит. ошибки Exception).")
        return jsonify(success=False, error="Внутренняя ошибка сервера."), 500

    if not files_to_proc:
         msg = "Нет файлов для обработки."
         processing_status[sess].update(status="Нет данных", error=msg, percent=100)
         logging.warning(f"({sess}) {msg}")
         logging.info(f"({sess}) Проверка существования рабочей папки (нет файлов для обработки): {work_dir}")
         path_exists = os.path.exists(work_dir)
         logging.info(f"({sess}) Результат проверки os.path.exists({work_dir}) (нет файлов для обработки): {path_exists}")
         if path_exists:
             logging.info(f"({sess}) Попытка удаления рабочей папки (нет файлов для обработки): {work_dir}")
             try:
                 shutil.rmtree(work_dir)
                 logging.info(f"({sess}) Рабочая папка {work_dir} успешно удалена (нет файлов для обработки).")
             except Exception as e_rm:
                 logging.error(f"({sess}) Ошибка при удалении рабочей папки {work_dir} (нет файлов для обработки): {e_rm}")
         else:
            logging.info(f"({sess}) Рабочая папка {work_dir} не найдена для удаления (нет файлов для обработки).")
         return jsonify(success=False, error=msg), 400

    # ───── Этап обработки файлов ─────
    results, empty, fails = [], [], []
    ref_path = None
    if smeta_type == "Смета ру": ref_path = REF_SMETA_RU
    elif smeta_type.startswith("Турбосметчик-"): ref_path = REF_TURBO
    elif smeta_type == "Грандсмета": ref_path = REF_GRAND
    else: logging.info(f"({sess}) Неизвестный тип сметы '{smeta_type}', реф. ширины не будут применены.")

    widths = None
    if ref_path and os.path.exists(ref_path):
        try:
            wb_r = openpyxl.load_workbook(ref_path, data_only=True)
            ws_r = wb_r.active
            widths = [ws_r.column_dimensions[get_column_letter(i)].width or 8.43 for i in range(1, 7)]
            wb_r.close()
            logging.info(f"({sess}) Успешно прочитаны реф. ширины из {os.path.basename(ref_path)}: {widths}")
        except Exception as e:
            logging.warning(f"({sess}) Не удалось прочитать реф. файл {os.path.basename(ref_path)}: {e}")
            widths = None
    elif ref_path: logging.warning(f"({sess}) Референсный файл не найден: {ref_path}")

    headers_common = None
    total_files_to_process = processing_status[sess]["total"]

    for idx, info in enumerate(files_to_proc, 1):
        file_disk_path = info["path"]
        # file_original_id теперь содержит скорректированный display_name
        file_original_id = info["original"]

        processing_status[sess].update(
            status=f"Обработка {idx}/{total_files_to_process}: {file_original_id}...", # Используем display_name
            percent=calc_percent("processing", processing_status[sess]["processed"], total_files_to_process),
        )
        logging.info(f"({sess}) Начало обработки файла {idx}/{total_files_to_process}: {file_original_id} (путь: {file_disk_path})")

        try:
            headers, rows = dispatcher.run_processor(smeta_type, file_disk_path)
            if not rows:
                logging.warning(f"({sess}) Файл '{file_original_id}' пустой или не содержит данных.")
                empty.append(file_original_id) # Добавляем display_name в список пустых
                # --- НОВОЕ: Обновляем статус немедленно ---
                processing_status[sess].update(status=f"⚠️ Файл пустой: {file_original_id}")
                # --- КОНЕЦ НОВОГО ---
            else:
                logging.info(f"({sess}) Файл '{file_original_id}' успешно обработан, строк: {len(rows)}.")
                results.append((file_original_id, headers, rows)) # Сохраняем display_name
                if headers_common is None and headers: headers_common = headers
                processing_status[sess]["processed"] += 1
        except Exception as proc_err:
            logging.exception(f"({sess}) Ошибка при обработке файла '{file_original_id}'")
            fails.append(file_original_id) # Добавляем display_name в список ошибок
            # --- НОВОЕ: Обновляем статус немедленно ---
            # Можно добавить детали ошибки, но пока просто факт ошибки
            processing_status[sess].update(status=f"❌ Ошибка файла: {file_original_id}")
            # --- КОНЕЦ НОВОГО ---
            # processing_status[sess]["error"] = f"Ошибка обработки {file_original_id}" # Старую строку можно убрать или оставить для логов/итогов

        processing_status[sess]["percent"] = calc_percent(
            "processing", processing_status[sess]["processed"], total_files_to_process
        )

    if not results:
        error_msg = "Не найдено данных для итогового файла."
        if fails: error_msg = f"Все {len(fails)} файла(ов) не удалось обработать."
        elif empty: error_msg = f"Все {len(empty)} файла(ов) оказались пустыми."
        processing_status[sess].update(status="Нет данных", error=error_msg, percent=100)
        logging.error(f"({sess}) {error_msg}")
        logging.info(f"({sess}) Проверка существования рабочей папки (нет данных для итогового файла): {work_dir}")
        path_exists = os.path.exists(work_dir)
        logging.info(f"({sess}) Результат проверки os.path.exists({work_dir}) (нет данных для итог. файла): {path_exists}")
        if path_exists:
            logging.info(f"({sess}) Попытка удаления рабочей папки (нет данных для итогового файла): {work_dir}")
            try:
                shutil.rmtree(work_dir)
                logging.info(f"({sess}) Рабочая папка {work_dir} успешно удалена (нет данных для итогового файла).")
            except Exception as e_rm:
                logging.error(f"({sess}) Ошибка при удалении рабочей папки {work_dir} (нет данных для итогового файла): {e_rm}")
        else:
            logging.info(f"({sess}) Рабочая папка {work_dir} не найдена для удаления (нет данных для итог. файла).")
        # Возвращаем списки ошибок/пустых
        return jsonify(success=False, error=error_msg, empty_files=empty, failed_files=fails), 400

    # ───── Формирование итогового Excel файла ─────
    processing_status[sess].update( status="Формирование итогового файла...", percent=calc_percent("finalize", 0, 0))
    logging.info(f"({sess}) Начало формирования итогового Excel файла...")

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = (re.sub(r'[^\w.-]+', '_', smeta_type).strip() or "Результат")[:31]
    if headers_common: ws.append(headers_common)

    # Добавляем данные из каждого успешно обработанного файла
    for file_original_id, headers, rows in results:
        # Добавляем разделитель с display_name, если файлов больше одного
        if len(results) > 1:
            r_idx = ws.max_row + 1
            max_col_letter = get_column_letter(max(ws.max_column, 6))
            merge_range_str = f'A{r_idx}:{max_col_letter}{r_idx}'
            try:
                 ws.merge_cells(merge_range_str)
                 cell = ws.cell(row=r_idx, column=1, value=file_original_id) # Используем display_name
                 cell.alignment = Alignment(horizontal="center", vertical="center")
                 cell.font = Font(bold=True)
            except Exception as merge_err:
                 ws.cell(row=r_idx, column=1, value=file_original_id).font = Font(bold=True)
                 logging.warning(f"({sess}) Не удалось объединить ячейки для разделителя '{file_original_id}': {merge_err}")

        # Добавляем строки данных из текущего файла
        row_count_before = ws.max_row
        for row_data in rows:
            if not row_data: continue
            current_row_idx = ws.max_row + 1
            if row_data[0] == "__FOOTER__":
                footer_text = row_data[1] if len(row_data) > 1 else ""
                footer_value = row_data[2] if len(row_data) > 2 else None
                try:
                    ws.merge_cells(start_row=current_row_idx, start_column=4, end_row=current_row_idx, end_column=5)
                    ws.cell(current_row_idx, 4, footer_text)
                    if footer_value is not None: ws.cell(current_row_idx, 6, footer_value)
                except Exception as footer_err: logging.warning(f"({sess}) Ошибка обработки __FOOTER__ в строке {current_row_idx}: {footer_err}")
                continue
            for col_idx, cell_value in enumerate(row_data, 1):
                if cell_value is not None: ws.cell(current_row_idx, col_idx, cell_value)
            if len(row_data) > 2 and row_data[1] is not None and row_data[2] is None and row_data[0] != '__FOOTER__':
                 try: ws.merge_cells(start_row=current_row_idx, start_column=2, end_row=current_row_idx, end_column=3)
                 except Exception as merge_err: logging.warning(f"({sess}) Не удалось объединить B-C в строке {current_row_idx}: {merge_err}")
        logging.debug(f"({sess}) Добавлено {ws.max_row - row_count_before} строк из '{file_original_id}'")

    # Применение форматирования и ширин
    try:
        if widths: apply_reference_widths(ws, widths)
        else: auto_adjust_column_width(ws)
        apply_formatting(ws)
        logging.info(f"({sess}) Форматирование применено.")
    except Exception as fmt_err: logging.error(f"({sess}) Ошибка применения форматирования: {fmt_err}")

    # Сохранение итогового файла
    try:
        safe_type_name = re.sub(r'[^\w.-]+', '_', smeta_type).strip() or "result"
        output_filename = f"{safe_type_name}_{sess[:8]}_processed.xlsx"
        output_filepath = os.path.join(RESULTS_FOLDER, output_filename)
        wb.save(output_filepath)
        wb.close()
        logging.info(f"({sess}) Итоговый файл сохранен как: {output_filepath}")
    except Exception as save_err:
        processing_status[sess].update(status="Ошибка сохранения", error=f"Не удалось сохранить итоговый файл: {save_err}", percent=100)
        logging.exception(f"({sess}) Не удалось сохранить итоговый файл")
        logging.info(f"({sess}) Проверка существования рабочей папки (ошибка сохранения итогового файла): {work_dir}")
        path_exists = os.path.exists(work_dir)
        logging.info(f"({sess}) Результат проверки os.path.exists({work_dir}) (ошибка сохранения итог. файла): {path_exists}")
        if path_exists:
            logging.info(f"({sess}) Попытка удаления рабочей папки (ошибка сохранения итогового файла): {work_dir}")
            try:
                shutil.rmtree(work_dir)
                logging.info(f"({sess}) Рабочая папка {work_dir} успешно удалена (ошибка сохранения итогового файла).")
            except Exception as e_rm:
                logging.error(f"({sess}) Ошибка при удалении рабочей папки {work_dir} (ошибка сохранения итогового файла): {e_rm}")
        else:
            logging.info(f"({sess}) Рабочая папка {work_dir} не найдена для удаления (ошибка сохранения итог. файла).")
        return jsonify(success=False, error="Ошибка сохранения итогового файла."), 500

    # Финальный статус и ответ клиенту
    processing_status[sess].update( status="Готово", percent=calc_percent("done", 0, 0), error=None )
    info_parts = []
    if results: info_parts.append(f"✅ Обработано: {len(results)}")
    if empty:   info_parts.append(f"⚠️ Пустые: {len(empty)}")
    if fails:   info_parts.append(f"❌ Ошибки: {len(fails)}")
    final_message = "; ".join(info_parts) if info_parts else "Обработка завершена, но нет данных для отображения."
    logging.info(f"({sess}) Обработка успешно завершена. {final_message}")
    logging.info(f"({sess}) Проверка существования рабочей папки (успешное завершение): {work_dir}")
    path_exists = os.path.exists(work_dir)
    logging.info(f"({sess}) Результат проверки os.path.exists({work_dir}) (успешное завершение): {path_exists}")
    if path_exists:
        logging.info(f"({sess}) Попытка удаления рабочей папки (успешное завершение): {work_dir}")
        try:
            shutil.rmtree(work_dir)
            logging.info(f"({sess}) Рабочая папка {work_dir} успешно удалена (успешное завершение).")
        except Exception as e_rm:
            logging.error(f"({sess}) Ошибка при удалении рабочей папки {work_dir} (успешное завершение): {e_rm}")
    else:
        logging.info(f"({sess}) Рабочая папка {work_dir} не найдена для удаления (успешное завершение).")

    # Возвращаем результат, включая списки пустых/ошибок
    return jsonify(
        success=True,
        message=final_message,
        download_url=url_for("download_file", filename=output_filename),
        download_filename=output_filename,
        empty_files=empty,  # Список display_name пустых файлов
        failed_files=fails  # Список display_name файлов с ошибками
    )

# ────────── progress & download ──────────
@app.route("/progress/<session_id>")
def progress(session_id):
    """Возвращает статус обработки для указанной сессии."""
    status = processing_status.get(session_id, {
        "status": "Инициализация", "processed": 0, "total": 0, "percent": 0, "error": "ID сессии не найден на сервере."
    })
    return jsonify(status)

@app.route("/download/<filename>")
def download_file(filename):
    """Отдает на скачивание итоговый файл."""
    if ".." in filename or filename.startswith(("/", "\\")):
        logging.warning(f"[БЕЗОПАСНОСТЬ] Попытка скачивания файла с небезопасным путем: {filename}")
        return "Недопустимое имя файла", 400

    file_path = os.path.join(RESULTS_FOLDER, filename)
    if not os.path.exists(file_path):
        logging.error(f"Запрошенный для скачивания файл не найден: {file_path}")
        return "Файл не найден", 404

    logging.info(f"Отправка файла для скачивания: {filename}")
    try:
        return send_from_directory(RESULTS_FOLDER, filename, as_attachment=True)
    except Exception as e:
        logging.exception(f"Ошибка при отправке файла '{filename}'")
        return "Ошибка сервера при отправке файла", 500

# --- Точка входа ---
if __name__ == "__main__":
    app.run(host="0.0.0.0", port=5000, debug=True)